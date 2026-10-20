"""
Готовит книгу sheets/psp_dashboards.xlsx для импорта в Google Sheets.

Python только выгружает данные на листы data_*. Все показатели, фильтры,
условное форматирование и графики на листах «Сверка», «Балансы» и
«Платёжный отчёт» считаются формулами самой таблицы (SUMIFS, COUNTIFS, IFERROR),
поэтому после импорта в Google Sheets дашборд живой: меняете фильтр или порог —
пересчитывается всё.

Запуск: python src/build_sheets.py
"""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "sheets"
OUT.mkdir(exist_ok=True)

PSP = ["PSP Alfa", "PSP Beta", "PSP Gamma"]
REGIONS = ["Москва", "Санкт-Петербург", "Другие регионы"]
COLORS = {"PSP Alfa": "2F6FDE", "PSP Beta": "8B5CF6", "PSP Gamma": "F08A24"}

FONT = "Arial"
H1 = Font(name=FONT, size=16, bold=True)
H2 = Font(name=FONT, size=12, bold=True)
BOLD = Font(name=FONT, size=10, bold=True)
BASE = Font(name=FONT, size=10)
MUTED = Font(name=FONT, size=9, color="6B7280")
INPUT = Font(name=FONT, size=10, color="0000FF", bold=True)
HEAD_FILL = PatternFill("solid", fgColor="F1F3F6")
INPUT_FILL = PatternFill("solid", fgColor="FFF4C2")
KPI_FILL = PatternFill("solid", fgColor="F7F8FA")
RED_FILL = PatternFill("solid", fgColor="FBE3E3")
LINE = Side(style="thin", color="E5E7EB")
BOTTOM = Border(bottom=LINE)

RUB = '#,##0 "₽";-#,##0 "₽";"–"'
INT = "#,##0"
PCT1 = "0.0%"
PCT2 = "0.00%"

# ---------- данные ----------
ledger = pd.read_csv(DATA / "ledger.csv")
psp_report = pd.read_csv(DATA / "psp_report.csv")
balances = pd.read_csv(DATA / "balances.csv")
payments = pd.read_csv(DATA / "payments_weekly.csv")
breaks_log = pd.read_csv(DATA / "breaks_log.csv")

con = sqlite3.connect(":memory:")
ledger.to_sql("ledger", con, index=False)
psp_report.to_sql("psp_report", con, index=False)
recon = pd.read_sql((ROOT / "sql" / "reconciliation.sql").read_text(encoding="utf-8"), con)
recon["is_break"] = (recon.recon_result != "matched").astype(int)
daily = (recon.groupby(["txn_date", "psp"])
         .agg(transactions=("txn_id", "size"), breaks=("is_break", "sum"))
         .reset_index())
DAYS = sorted(daily.txn_date.unique())
for df, col in ((daily, "txn_date"), (balances, "balance_date"), (breaks_log, "txn_date")):
    df[col] = pd.to_datetime(df[col])

wb = Workbook()


def data_sheet(name, df, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in df.itertuples(index=False):
        ws.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in row])
    for c in ws["A"][1:] + ws["B"][1:]:
        if c.is_date:
            c.number_format = "DD.MM.YYYY"
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "BFC5CE"
    return ws, len(df) + 1


d_daily, n_daily = data_sheet(
    "data_сверка", daily[["txn_date", "psp", "transactions", "breaks"]],
    ["Дата", "PSP", "Транзакций", "Расхождений"])
d_log, n_log = data_sheet(
    "data_журнал", breaks_log[["break_id", "txn_date", "psp", "txn_id", "break_type", "amount_rub", "reason", "status"]],
    ["Номер", "Дата", "PSP", "ID транзакции", "Тип", "Сумма, ₽", "Причина", "Статус"])
d_log.column_dimensions["G"].width = 52
d_bal, n_bal = data_sheet(
    "data_балансы", balances[["balance_date", "psp", "expected_balance_rub", "actual_balance_rub"]],
    ["Дата", "PSP", "Ожидаемый баланс, ₽", "Фактический баланс, ₽"])
pay = payments.assign(type=payments.type.map({"deposit": "Депозит", "withdrawal": "Вывод"}))
d_pay, n_pay = data_sheet(
    "data_платежи", pay[["week", "psp", "region", "type", "transactions", "successful", "failed", "volume_rub", "fee_rub"]],
    ["Неделя", "PSP", "Регион", "Тип", "Транзакций", "Успешных", "Неуспешных", "Оборот, ₽", "Комиссия, ₽"])
for ws, cols in ((d_log, "F"), (d_bal, "CD"), (d_pay, "HI")):
    for col in cols:
        for c in ws[col][1:]:
            c.number_format = RUB


def rng(ws_name, col, n):
    return f"'{ws_name}'!${col}$2:${col}${n}"


DD = {k: rng("data_сверка", c, n_daily) for k, c in zip(["date", "psp", "tx", "br"], "ABCD")}
DL = {k: rng("data_журнал", c, n_log) for k, c in zip(["id", "date", "psp", "txn", "type", "amt", "reason", "status"], "ABCDEFGH")}
DB = {k: rng("data_балансы", c, n_bal) for k, c in zip(["date", "psp", "exp", "act"], "ABCD")}
DP = {k: rng("data_платежи", c, n_pay) for k, c in zip(["week", "psp", "region", "type", "tx", "ok", "fail", "vol", "fee"], "ABCDEFGHI")}


def head(ws, row, col, values, widths=None):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, BOTTOM
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)


def kpi(ws, row, col, label, formula, fmt):
    a = ws.cell(row, col, label)
    a.font, a.fill = MUTED, KPI_FILL
    b = ws.cell(row + 1, col, formula)
    b.font, b.fill, b.number_format = Font(name=FONT, size=18, bold=True), KPI_FILL, fmt


def style_body(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if not cell.font.bold:
                cell.font = BASE
            cell.border = BOTTOM


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def dropdown(ws, cell, options, label_cell, label):
    ws[label_cell] = label
    ws[label_cell].font = MUTED
    ws[cell] = options[0]
    ws[cell].font, ws[cell].fill = INPUT, INPUT_FILL
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws[cell])


def crit(cell):
    """«Все» в фильтре превращается в шаблон «*», который SUMIFS понимает как любое значение."""
    return f'IF({cell}="Все","*",{cell})'


# ---------- лист «Сверка» ----------
ws = wb.active
ws.title = "Сверка"
ws.sheet_view.showGridLines = False
widths(ws, [16, 14, 14, 14, 14, 3, 16, 12, 12, 12, 12, 12, 12])
ws["A1"] = "Ежедневная сверка учёта с отчётами PSP"
ws["A1"].font = H1
ws["A2"] = "Сентябрь 2026 · 3 PSP · собственный проект на синтетических данных, суммы в рублях"
ws["A2"].font = MUTED
dropdown(ws, "B4", ["Все"] + PSP, "A4", "Фильтр PSP:")
ws["C4"] = "Порог match rate:"
ws["C4"].font = MUTED
ws["D4"] = 0.99
ws["D4"].font, ws["D4"].fill, ws["D4"].number_format = INPUT, INPUT_FILL, PCT1

f_psp = crit("$B$4")
kpi(ws, 6, 1, "Транзакций", f"=SUMIFS({DD['tx']},{DD['psp']},{f_psp})", INT)
kpi(ws, 6, 2, "Расхождений", f"=SUMIFS({DD['br']},{DD['psp']},{f_psp})", INT)
kpi(ws, 6, 3, "Match rate", "=IFERROR(1-B7/A7,0)", PCT2)
kpi(ws, 6, 4, "Открыто / в работе", f'=COUNTIFS({DL["psp"]},{f_psp},{DL["status"]},"<>Закрыто")', INT)
kpi(ws, 6, 5, "Сумма открытых", f'=SUMIFS({DL["amt"]},{DL["psp"]},{f_psp},{DL["status"]},"<>Закрыто")', RUB)

ws["A10"] = "Итоги по PSP"
ws["A10"].font = H2
head(ws, 11, 1, ["PSP", "Транзакций", "Расхождений", "Match rate", "Ниже порога, дней"])
for i, p in enumerate(PSP):
    r = 12 + i
    ws.cell(r, 1, p).font = Font(name=FONT, size=10, bold=True, color=COLORS[p])
    ws.cell(r, 2, f"=SUMIFS({DD['tx']},{DD['psp']},$A{r})").number_format = INT
    ws.cell(r, 3, f"=SUMIFS({DD['br']},{DD['psp']},$A{r})").number_format = INT
    ws.cell(r, 4, f"=IFERROR(1-C{r}/B{r},0)").number_format = PCT2
    col = get_column_letter(8 + i)
    ws.cell(r, 5, f"=COUNTIF(${col}$19:${col}${18 + len(DAYS)},\"<\"&$D$4)").number_format = INT
style_body(ws, 12, 14, 1, 5)
ws.conditional_formatting.add("D12:D14", CellIsRule(operator="lessThan", formula=["$D$4"], fill=RED_FILL))

ws["A16"] = "Расхождения по типам"
ws["A16"].font = H2
TYPES = ["Статус", "Сумма", "Комиссия", "Нет у PSP", "Нет в учёте"]
head(ws, 17, 1, ["Тип"] + PSP)
for i, t in enumerate(TYPES):
    r = 18 + i
    ws.cell(r, 1, t)
    for j, p in enumerate(PSP):
        ws.cell(r, 2 + j, f'=COUNTIFS({DL["type"]},$A{r},{DL["psp"]},"{p}")').number_format = INT
style_body(ws, 18, 22, 1, 4)

ws["G16"] = "Match rate по дням"
ws["G16"].font = H2
head(ws, 18, 7, ["Дата"] + PSP)
for k, d in enumerate(DAYS):
    r = 19 + k
    ws.cell(r, 7, pd.Timestamp(d).to_pydatetime()).number_format = "DD.MM"
    for j, p in enumerate(PSP):
        ws.cell(r, 8 + j,
                f'=IFERROR(1-SUMIFS({DD["br"]},{DD["date"]},$G{r},{DD["psp"]},"{p}")'
                f'/SUMIFS({DD["tx"]},{DD["date"]},$G{r},{DD["psp"]},"{p}"),"")').number_format = PCT1
last = 18 + len(DAYS)
style_body(ws, 19, last, 7, 10)
ws.conditional_formatting.add(f"H19:J{last}", CellIsRule(operator="lessThan", formula=["$D$4"], fill=RED_FILL))

ch = LineChart()
ch.title, ch.height, ch.width = "Match rate по дням", 8, 18
ch.y_axis.number_format, ch.y_axis.scaling.min, ch.y_axis.scaling.max = "0%", 0.9, 1.0
ch.add_data(Reference(ws, min_col=8, max_col=10, min_row=18, max_row=last), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=7, min_row=19, max_row=last))
for s, p in zip(ch.series, PSP):
    s.graphicalProperties.line.solidFill, s.graphicalProperties.line.width, s.smooth = COLORS[p], 22000, False
ws.add_chart(ch, "L4")

bc = BarChart()
bc.type, bc.grouping, bc.overlap = "col", "stacked", 100
bc.title, bc.height, bc.width = "Расхождения по PSP и типам", 8, 18
bc.add_data(Reference(ws, min_col=1, max_col=4, min_row=17, max_row=22), titles_from_data=True, from_rows=True)
bc.set_categories(Reference(ws, min_col=2, max_col=4, min_row=17))
ws.add_chart(bc, "L22")

ws["A25"] = "Как читать: желтые ячейки — фильтр и порог, их можно менять. Красным подсвечены дни и PSP ниже порога."
ws["A25"].font = MUTED

# ---------- лист «Журнал расхождений» ----------
wj = wb.create_sheet("Журнал")
wj.sheet_view.showGridLines = False
widths(wj, [10, 12, 12, 17, 13, 13, 54, 12])
wj["A1"] = "Журнал расхождений"
wj["A1"].font = H1
wj["A2"] = "Каждое расхождение проходит путь «обнаружено → причина → закрыто». Статус и причину заполняет аналитик."
wj["A2"].font = MUTED
dropdown(wj, "B4", ["Все", "Открыто", "В работе", "Закрыто"], "A4", "Статус:")
head(wj, 6, 1, ["Номер", "Дата", "PSP", "ID транзакции", "Тип", "Сумма, ₽", "Причина", "Статус"])
for k in range(n_log - 1):
    r, src = 7 + k, 2 + k
    show = f'OR($B$4="Все",\'data_журнал\'!$H${src}=$B$4)'
    for j, col in enumerate("ABCDEFGH"):
        c = wj.cell(r, 1 + j, f'=IF({show},\'data_журнал\'!{col}{src},"")')
        if col == "F":
            c.number_format = RUB
style_body(wj, 7, 6 + n_log - 1, 1, 8)
wj.conditional_formatting.add(f"A7:H{6 + n_log - 1}",
                               FormulaRule(formula=['$H7="Открыто"'], fill=RED_FILL))
wj.conditional_formatting.add(f"A7:H{6 + n_log - 1}",
                               FormulaRule(formula=['$H7="В работе"'], fill=INPUT_FILL))
wj.freeze_panes = "A7"

# ---------- лист «Балансы» ----------
wbal = wb.create_sheet("Балансы")
wbal.sheet_view.showGridLines = False
widths(wbal, [12, 14, 14, 14, 3, 18, 14, 14, 14])
wbal["A1"] = "Контроль балансов на счетах PSP"
wbal["A1"].font = H1
wbal["A2"] = "Дельта = фактический баланс из выписки PSP − ожидаемый по учёту (остаток + депозиты − выводы − комиссии)"
wbal["A2"].font = MUTED
wbal["A4"] = "Порог эскалации, ₽:"
wbal["A4"].font = MUTED
wbal["B4"] = 50000
wbal["B4"].font, wbal["B4"].fill, wbal["B4"].number_format = INPUT, INPUT_FILL, RUB
wbal["C4"] = "Окно тренда, дней:"
wbal["C4"].font = MUTED
wbal["D4"] = 7
wbal["D4"].font, wbal["D4"].fill = INPUT, INPUT_FILL

head(wbal, 6, 1, ["Дата"] + [f"{p}, дельта" for p in PSP])
for k, d in enumerate(DAYS):
    r = 7 + k
    wbal.cell(r, 1, pd.Timestamp(d).to_pydatetime()).number_format = "DD.MM"
    for j, p in enumerate(PSP):
        key = f"$A{r}"
        wbal.cell(r, 2 + j,
                  f'=SUMIFS({DB["act"]},{DB["date"]},{key},{DB["psp"]},"{p}")'
                  f'-SUMIFS({DB["exp"]},{DB["date"]},{key},{DB["psp"]},"{p}")').number_format = RUB
blast = 6 + len(DAYS)
style_body(wbal, 7, blast, 1, 4)
wbal.conditional_formatting.add(f"B7:D{blast}", FormulaRule(formula=["ABS(B7)>$B$4"], fill=RED_FILL))

head(wbal, 6, 6, ["Сводка", *PSP])
rows = [
    ("Дельта на конец месяца", lambda c: f"={c}{blast}", RUB),
    ("Минимум дельты", lambda c: f"=MIN({c}7:{c}{blast})", RUB),
    ("Дней выше порога", lambda c: f'=COUNTIF({c}7:{c}{blast},">"&$B$4)+COUNTIF({c}7:{c}{blast},"<"&-$B$4)', INT),
    ("Дней, когда дельта ниже вчерашней", lambda c: f"=SUMPRODUCT(--({c}8:{c}{blast}<{c}7:{c}{blast - 1}))", INT),
]
for i, (label, f, fmt) in enumerate(rows):
    r = 7 + i
    wbal.cell(r, 6, label)
    for j in range(3):
        wbal.cell(r, 7 + j, f(get_column_letter(2 + j))).number_format = fmt
style_body(wbal, 7, 10, 6, 9)

# тренд: сравниваем среднюю дельту за последнее окно с предыдущим окном.
# Разовый провал (задержка зачисления) тренд не меняет, систематическая недоплата — меняет.
head(wbal, 12, 6, ["Тренд", *PSP])
wbal.cell(13, 6, "Средняя дельта, последнее окно")
wbal.cell(14, 6, "Средняя дельта, предыдущее окно")
wbal.cell(15, 6, "Сигнал")
for j in range(3):
    c = get_column_letter(2 + j)
    wbal.cell(13, 7 + j, f"=AVERAGE(OFFSET({c}{blast},-$D$4+1,0,$D$4,1))").number_format = RUB
    wbal.cell(14, 7 + j, f"=AVERAGE(OFFSET({c}{blast},-2*$D$4+1,0,$D$4,1))").number_format = RUB
    col = get_column_letter(7 + j)
    wbal.cell(15, 7 + j, f'=IF(AND({col}13<{col}14,ABS({col}13)>$B$4/2),"Проверить","Норма")')
style_body(wbal, 13, 15, 6, 9)
wbal.conditional_formatting.add("G15:I15", CellIsRule(operator="equal", formula=['"Проверить"'], fill=RED_FILL))

lc = LineChart()
lc.title, lc.height, lc.width = "Дельта баланса по дням, ₽", 9, 20
lc.y_axis.number_format = "#,##0"
lc.add_data(Reference(wbal, min_col=2, max_col=4, min_row=6, max_row=blast), titles_from_data=True)
lc.set_categories(Reference(wbal, min_col=1, min_row=7, max_row=blast))
for s, p in zip(lc.series, PSP):
    s.graphicalProperties.line.solidFill, s.graphicalProperties.line.width, s.smooth = COLORS[p], 22000, False
wbal.add_chart(lc, "F18")

# ---------- лист «Платёжный отчёт» ----------
wp = wb.create_sheet("Платёжный отчёт")
wp.sheet_view.showGridLines = False
widths(wp, [20, 15, 16, 15, 15, 15, 3, 18, 14])
wp["A1"] = "Платёжный отчёт для трейдеров и финансов"
wp["A1"].font = H1
wp["A2"] = "Сентябрь 2026 · онлайн-букмекер · 3 PSP · Москва, Санкт-Петербург, другие регионы"
wp["A2"].font = MUTED
dropdown(wp, "B4", ["Все"] + PSP, "A4", "PSP:")
dropdown(wp, "B5", ["Все"] + REGIONS, "A5", "Регион:")
dropdown(wp, "B6", ["Все", "Депозит", "Вывод"], "A6", "Тип операции:")

flt = lambda extra="": (f'{DP["psp"]},{crit("$B$4")},{DP["region"]},{crit("$B$5")},{DP["type"]},{crit("$B$6")}{extra}')
kpi(wp, 8, 1, "Оборот", f"=SUMIFS({DP['vol']},{flt()})", RUB)
kpi(wp, 8, 2, "Транзакций", f"=SUMIFS({DP['tx']},{flt()})", INT)
kpi(wp, 8, 3, "Доля успешных", f"=IFERROR(SUMIFS({DP['ok']},{flt()})/B9,0)", PCT1)
kpi(wp, 8, 4, "Комиссии", f"=SUMIFS({DP['fee']},{flt()})", RUB)
kpi(wp, 8, 5, "Эффективная комиссия", "=IFERROR(D9/A9,0)", PCT2)

wp["A12"] = "По провайдерам"
wp["A12"].font = H2
head(wp, 13, 1, ["PSP", "Доля успешных", "Оборот, ₽", "Комиссии, ₽", "Эффективная комиссия", "Отклонено, шт."])
for i, p in enumerate(PSP):
    r = 14 + i
    f = f'{DP["psp"]},$A{r},{DP["region"]},{crit("$B$5")},{DP["type"]},{crit("$B$6")}'
    wp.cell(r, 1, p).font = Font(name=FONT, size=10, bold=True, color=COLORS[p])
    wp.cell(r, 2, f"=IFERROR(SUMIFS({DP['ok']},{f})/SUMIFS({DP['tx']},{f}),0)").number_format = PCT1
    wp.cell(r, 3, f"=SUMIFS({DP['vol']},{f})").number_format = RUB
    wp.cell(r, 4, f"=SUMIFS({DP['fee']},{f})").number_format = RUB
    wp.cell(r, 5, f"=IFERROR(D{r}/C{r},0)").number_format = PCT2
    wp.cell(r, 6, f"=SUMIFS({DP['fail']},{f})").number_format = INT
style_body(wp, 14, 16, 1, 6)
wp.conditional_formatting.add("E14:E16", FormulaRule(formula=["E14=MAX($E$14:$E$16)"], fill=RED_FILL))

wp["A19"] = "По регионам"
wp["A19"].font = H2
head(wp, 20, 1, ["Регион", "Доля успешных", "Оборот, ₽", "Транзакций"])
for i, reg in enumerate(REGIONS):
    r = 21 + i
    f = f'{DP["region"]},$A{r},{DP["psp"]},{crit("$B$4")},{DP["type"]},{crit("$B$6")}'
    wp.cell(r, 1, reg)
    wp.cell(r, 2, f"=IFERROR(SUMIFS({DP['ok']},{f})/SUMIFS({DP['tx']},{f}),0)").number_format = PCT1
    wp.cell(r, 3, f"=SUMIFS({DP['vol']},{f})").number_format = RUB
    wp.cell(r, 4, f"=SUMIFS({DP['tx']},{f})").number_format = INT
style_body(wp, 21, 23, 1, 4)

wp["A26"] = "По неделям"
wp["A26"].font = H2
WEEKS = sorted(payments.week.unique())
head(wp, 27, 1, ["Неделя", "Депозиты, ₽", "Выводы, ₽", "Доля успешных"])
for i, w in enumerate(WEEKS):
    r = 28 + i
    base = f'{DP["week"]},$A{r},{DP["psp"]},{crit("$B$4")},{DP["region"]},{crit("$B$5")}'
    wp.cell(r, 1, w)
    wp.cell(r, 2, f'=SUMIFS({DP["vol"]},{base},{DP["type"]},"Депозит")').number_format = RUB
    wp.cell(r, 3, f'=SUMIFS({DP["vol"]},{base},{DP["type"]},"Вывод")').number_format = RUB
    wp.cell(r, 4, f"=IFERROR(SUMIFS({DP['ok']},{base},{DP['type']},{crit('$B$6')})"
                  f"/SUMIFS({DP['tx']},{base},{DP['type']},{crit('$B$6')}),0)").number_format = PCT1
style_body(wp, 28, 27 + len(WEEKS), 1, 4)

b1 = BarChart()
b1.type, b1.title, b1.height, b1.width = "col", "Оборот по PSP, ₽", 8, 15
b1.add_data(Reference(wp, min_col=3, min_row=13, max_row=16), titles_from_data=True)
b1.set_categories(Reference(wp, min_col=1, min_row=14, max_row=16))
b1.legend = None
b1.y_axis.number_format = "#,##0"
wp.add_chart(b1, "H4")

b2 = BarChart()
b2.type, b2.title, b2.height, b2.width = "bar", "Доля успешных платежей по регионам", 7, 15
b2.add_data(Reference(wp, min_col=2, min_row=20, max_row=23), titles_from_data=True)
b2.set_categories(Reference(wp, min_col=1, min_row=21, max_row=23))
b2.legend = None
b2.x_axis.number_format = "0%"
b2.y_axis.scaling.min, b2.y_axis.scaling.max = 0.8, 0.95
wp.add_chart(b2, "H21")

wp["A34"] = ("Как читать: желтые ячейки — фильтры. Отклонённый платёж не равен потере денег: часть отказов "
             "на стороне игрока (нет средств, отмена). Сравнивать провайдеров корректно при одинаковом наборе регионов.")
wp["A34"].font = MUTED

# порядок листов: дашборды впереди, данные в конце
wb._sheets = [wb["Сверка"], wb["Журнал"], wb["Балансы"], wb["Платёжный отчёт"],
              d_daily, d_log, d_bal, d_pay]
for s in wb.worksheets:
    s.sheet_view.zoomScale = 100

out = OUT / "psp_dashboards.xlsx"
wb.save(out)
print("written", out)
